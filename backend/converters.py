"""Convert PDF / DOCX / XLSX / image bytes into text chunks the embedding model can use.

All functions operate on raw bytes so the same code path works for files already
sitting on disk (resource pool ingestion) and files uploaded in-request (ephemeral
question attachments) without needing a temp file.
"""
import io

import pymupdf as fitz
import openpyxl
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ollama_client import describe_image

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
XLSX_EXTENSIONS = {".xlsx", ".xls"}
PPTX_EXTENSIONS = {".pptx"}
TEXT_EXTENSIONS = {".txt", ".md"}

ALL_EXTENSIONS = (
    IMAGE_EXTENSIONS | PDF_EXTENSIONS | DOCX_EXTENSIONS | XLSX_EXTENSIONS | PPTX_EXTENSIONS | TEXT_EXTENSIONS
)


async def extract_pdf(data: bytes) -> list[dict]:
    """One chunk per page. Pages with no extractable text (likely scanned) are
    rendered to an image and described by the vision model instead."""
    chunks = []
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if text:
                chunks.append({"text": text, "extra": {"page": page_num}})
            else:
                pix = page.get_pixmap(dpi=150)
                png_bytes = pix.tobytes("png")
                description = await describe_image(
                    png_bytes, "Transcribe and describe all text and content visible on this page."
                )
                if description.strip():
                    chunks.append({"text": description, "extra": {"page": page_num, "ocr": True}})
    finally:
        doc.close()
    return chunks


def extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_xlsx(data: bytes) -> list[dict]:
    """One chunk per data row, formatted as 'Column: value | Column: value'.
    This naturally keeps e.g. Question/Answer/Year columns tied together."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    chunks = []
    try:
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(h).strip() if h is not None else "" for h in rows[0]]
            for row_num, row in enumerate(rows[1:], start=2):
                pairs = [
                    f"{header[i]}: {row[i]}"
                    for i in range(min(len(header), len(row)))
                    if row[i] is not None and str(row[i]).strip()
                ]
                if pairs:
                    chunks.append({"text": " | ".join(pairs), "extra": {"sheet": sheet.title, "row": row_num}})
    finally:
        wb.close()
    return chunks


async def extract_pptx(data: bytes) -> list[dict]:
    """One chunk per slide (title, body text, tables). Slides with little or no
    extractable text (e.g. a diagram-only slide) have their first embedded image
    described by the vision model instead."""
    prs = Presentation(io.BytesIO(data))
    chunks = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        parts = []
        images = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            elif shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    images.append(shape.image.blob)
                except Exception:
                    pass

        text = "\n".join(parts).strip()
        extra = {"slide": slide_num}
        if len(text) < 20 and images:
            description = await describe_image(
                images[0], "Describe this slide's image in detail, including any visible text or diagrams."
            )
            if description.strip():
                text = (text + "\n" + description).strip()
                extra["ocr"] = True
        if text:
            chunks.append({"text": text, "extra": extra})
    return chunks


async def extract_image(data: bytes) -> str:
    return await describe_image(data, "Describe this image in detail, including any visible text.")
